"""Build dist/JobList.xlsx from a merged roles JSON (output of tools/merge-roles.js).

Usage: python3 tools/make-xlsx.py roles.json
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

roles = json.loads(Path(sys.argv[1]).read_text())
out = Path(__file__).resolve().parent.parent / 'dist' / 'JobList.xlsx'

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', start_color='1F3A5F')
HEAD_FONT = Font(name=FONT, bold=True, color='FFFFFF')
BODY = Font(name=FONT, size=10)
LINK = Font(name=FONT, size=10, color='0563C1', underline='single')

COLS = [  # header, width
    ('Deadline', 12), ('Days left', 9), ('Company', 24), ('Role (click to open)', 46),
    ('Type', 11), ('Category', 26), ('Location', 24), ('Remote?', 10), ('Pay (as listed)', 22),
    ('Fit (1-5)', 8), ('Resume to use', 20), ('Why it fits you', 50), ('Notes / eligibility', 70),
    ('My status', 13), ('Date applied', 12),
]
STATUSES = ['To apply', 'In progress', 'Applied', 'Assessment', 'Interview R1', 'Interview R2',
            'Final round', 'Offer', 'Rejected', 'Ghosted', 'Withdrawn', 'Skip']


def as_date(s):
    try:
        y, m, d = (int(x) for x in str(s)[:10].split('-'))
        return date(y, m, d)
    except (ValueError, TypeError):
        return None


wb = Workbook()
ws = wb.active
ws.title = 'All Roles'
for c, (h, w) in enumerate(COLS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 30

for r, j in enumerate(roles, 2):
    d = as_date(j.get('deadline'))
    vals = [d or 'Rolling', None, j.get('company'), j.get('role'), j.get('type'), j.get('category'),
            j.get('location'), j.get('remote'), j.get('pay') or 'Not listed', j.get('fit'),
            j.get('resume'), j.get('why'), j.get('notes'), 'To apply', None]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = BODY
        cell.alignment = Alignment(vertical='top', wrap_text=c in (4, 12, 13))
    ws.cell(row=r, column=1).number_format = 'mmm d, yyyy'
    ws.cell(row=r, column=2, value=f'=IF(ISNUMBER(A{r}),A{r}-TODAY(),"")')
    ws.cell(row=r, column=15).number_format = 'mmm d, yyyy'
    if j.get('link'):
        cell = ws.cell(row=r, column=4)
        cell.hyperlink = j['link']
        cell.font = LINK

last = len(roles) + 1
ws.freeze_panes = 'D2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{last}'

dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'N2:N{last}')

rng = f'A2:{get_column_letter(len(COLS))}{last}'
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['OR($N2="Applied",$N2="Assessment",$N2="Interview R1",$N2="Interview R2",$N2="Final round",$N2="Offer")'],
    fill=PatternFill('solid', start_color='D9EAD3'), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['OR($N2="Rejected",$N2="Withdrawn",$N2="Skip",$N2="Ghosted")'],
    font=Font(name=FONT, color='9E9E9E'), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['AND(ISNUMBER($B2),$B2<0)'], font=Font(name=FONT, color='9E9E9E', strike=True), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['AND(ISNUMBER($B2),$B2<=7)'], fill=PatternFill('solid', start_color='F4CCCC')))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=['AND(ISNUMBER($B2),$B2<=21)'], fill=PatternFill('solid', start_color='FFF2CC')))
ws.conditional_formatting.add(f'J2:J{last}', FormulaRule(
    formula=['$J2>=4'], font=Font(name=FONT, bold=True, color='0B5132')))

# ---- Summary tab
s = wb.create_sheet('Summary', 0)
s['A1'] = f'Job list — {len(roles)} open roles (researched {date.today():%b %d, %Y})'
s['A1'].font = Font(name=FONT, size=14, bold=True)
notes = [
    'How to use: go to "All Roles". Use the filter arrows in the header row to filter by Category, Location, Type or Fit.',
    'Red rows are due within 7 days; yellow within 3 weeks; struck-through gray = deadline passed.',
    'Set "My status" (dropdown) as you apply; applied rows turn green.',
    'Found via web search: open each link to confirm the posting is live before applying. Eligibility notes are in the Notes column.',
]
for i, n in enumerate(notes, 2):
    s.cell(row=i, column=1, value=n).font = Font(name=FONT, size=10, italic=True)

hdr_row = 7
AR = "'All Roles'!"


def header(row, col, labels):
    for k, lab in enumerate(labels):
        cell = s.cell(row=row, column=col + k, value=lab)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL


header(hdr_row, 1, ['Category', 'Roles', 'Fit 4-5'])
cats = sorted({j.get('category') or 'Other' for j in roles})
for i, ccat in enumerate(cats, hdr_row + 1):
    s.cell(row=i, column=1, value=ccat).font = BODY
    s.cell(row=i, column=2, value=f'=COUNTIF({AR}$F$2:$F${last},A{i})').font = BODY
    s.cell(row=i, column=3, value=f'=COUNTIFS({AR}$F$2:$F${last},A{i},{AR}$J$2:$J${last},">=4")').font = BODY
tot = hdr_row + 1 + len(cats)
s.cell(row=tot, column=1, value='Total').font = Font(name=FONT, bold=True)
s.cell(row=tot, column=2, value=f'=SUM(B{hdr_row + 1}:B{tot - 1})').font = Font(name=FONT, bold=True)
s.cell(row=tot, column=3, value=f'=SUM(C{hdr_row + 1}:C{tot - 1})').font = Font(name=FONT, bold=True)

header(hdr_row, 5, ['Deadline window', 'Roles'])
windows = [('Due in 7 days', f'=COUNTIFS({AR}$B$2:$B${last},">=0",{AR}$B$2:$B${last},"<=7")'),
           ('Due in 8-30 days', f'=COUNTIFS({AR}$B$2:$B${last},">=8",{AR}$B$2:$B${last},"<=30")'),
           ('Due later', f'=COUNTIF({AR}$B$2:$B${last},">30")'),
           ('Rolling / no deadline', f'=COUNTIF({AR}$A$2:$A${last},"Rolling")'),
           ('Deadline passed', f'=COUNTIF({AR}$B$2:$B${last},"<0")')]
for i, (lab, f) in enumerate(windows, hdr_row + 1):
    s.cell(row=i, column=5, value=lab).font = BODY
    s.cell(row=i, column=6, value=f).font = BODY

header(hdr_row, 8, ['My progress', 'Count'])
prog = ['To apply', 'Applied', 'Interview R1', 'Offer', 'Skip']
for i, p in enumerate(prog, hdr_row + 1):
    s.cell(row=i, column=8, value=p).font = BODY
    s.cell(row=i, column=9, value=f'=COUNTIF({AR}$N$2:$N${last},H{i})').font = BODY

for col, w in zip('ABCDEFGHI', [38, 8, 8, 3, 22, 8, 3, 16, 8]):
    s.column_dimensions[col].width = w

wb.save(out)
print(f'wrote {out} ({len(roles)} roles)')
